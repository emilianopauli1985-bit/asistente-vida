import streamlit as st
import pandas as pd
from datetime import datetime
import io
from collections import Counter
import re
import string
import os

# Configuración inicial de la página
st.set_page_config(page_title="Asistente Operativo - VIDA", layout="wide")

# =========================================================
# VARIABLES DE MEMORIA (Para limpiar los archivos subidos)
# =========================================================
if 'reset_armador' not in st.session_state:
    st.session_state.reset_armador = 0
if 'reset_validador' not in st.session_state:
    st.session_state.reset_validador = 0
if 'reset_unificador' not in st.session_state:
    st.session_state.reset_unificador = 0
if 'reset_pdf' not in st.session_state:
    st.session_state.reset_pdf = 0

# =========================================================
# LISTAS DESPLEGABLES EXTRAÍDAS DE LA PLANTILLA OFICIAL VIDA
# =========================================================
LISTAS_DESPLEGABLES = {
    'Acción': ['Agregar', 'Eliminar', 'Actualizar'],
    '*Tipo Id.': ['D.N.I.', 'Pasaporte'],
    'Nro. de grupo': ['1-Titular', '2-Titular', '3-Titular', '4-Titular', '5-Titular', '6-Titular', '7-Titular']
}

# =========================================================
# FUNCIÓN PARA INYECTAR LAS LISTAS DESPLEGABLES EN EXCEL
# =========================================================
def inyectar_listas_desplegables(writer, df_export, workbook, sheet_name):
    worksheet_main = writer.sheets[sheet_name]
    
    # Verificamos si la hoja oculta DataValidation ya existe (útil para archivos Multi-Hoja)
    worksheet_dv = workbook.get_worksheet_by_name('DataValidation')
    
    if worksheet_dv is None:
        worksheet_dv = workbook.add_worksheet('DataValidation')
        col_idx = 0
        for col_name, valores in LISTAS_DESPLEGABLES.items():
            worksheet_dv.write(0, col_idx, col_name)
            worksheet_dv.write_column(1, col_idx, valores)
            col_idx += 1
        worksheet_dv.hide()
        
    col_idx = 0
    for col_name, valores in LISTAS_DESPLEGABLES.items():
        if col_name in df_export.columns:
            main_col_idx = list(df_export.columns).index(col_name)
            
            if col_idx < 26:
                letra_col_dv = string.ascii_uppercase[col_idx]
            else:
                letra_col_dv = string.ascii_uppercase[col_idx // 26 - 1] + string.ascii_uppercase[col_idx % 26]
                
            rango_formula = f"=DataValidation!${letra_col_dv}$2:${letra_col_dv}${len(valores)+1}"
            
            worksheet_main.data_validation(1, main_col_idx, 5000, main_col_idx, {
                'validate': 'list',
                'source': rango_formula
            })
            
        col_idx += 1

# =========================================================
# FUNCIONES DE LIMPIEZA Y PINTADO (ESTILOS)
# =========================================================
def procesar_dni_cuil(valor):
    if pd.isna(valor): return ""
    texto = str(valor).strip()
    if texto.lower() == 'nan' or not texto: return ""
    
    # 1. Limpiamos todos los caracteres que no sean números (guiones, puntos, letras)
    numeros = re.sub(r'[^0-9]', '', texto)
    
    # 2. Si tiene 11 números exactos y arranca con un prefijo de CUIL (20, 23, 24, 27, 30, 33)
    if len(numeros) == 11 and numeros[:2] in ['20', '23', '24', '27', '30', '33']:
        return numeros[2:10]
        
    # 3. Si tiene 10 números y arranca con prefijo (Ej: se olvidó el dígito verificador del final)
    elif len(numeros) == 10 and numeros[:2] in ['20', '23', '24', '27', '30', '33']:
        return numeros[2:]
        
    # 4. Si el usuario mandó "DNI-Digito" ej: 29452749-5 (un solo guion pero el DNI primero)
    if '-' in texto:
        partes = texto.split('-')
        if len(partes) == 2 and len(partes[0]) >= 7:
            return re.sub(r'[^0-9]', '', partes[0])
            
    # 5. En cualquier otro caso (ej: si mandó un DNI limpio), devolvemos los números.
    return numeros

def pintar_duplicados(s):
    dups = s.duplicated(keep=False) & (s != '')
    return ['background-color: #FFF2CC; color: #9C6500; font-weight: bold' if v else '' for v in dups]

# Fondo de pantalla y estilos de contraste visual
st.markdown(
    """
    <style>
    .stApp {
        background-image: linear-gradient(rgba(0, 0, 0, 0.6), rgba(0, 0, 0, 0.6)), url("https://raw.githubusercontent.com/emilianopauli1985-bit/validador-ap/main/istockphoto-2046622743-612x612.jpg");
        background-size: cover;
        background-position: center;
        background-attachment: fixed;
    }
    .stAlert, [data-testid="stDataFrame"] {
        background-color: rgba(255, 255, 255, 0.95) !important;
        border-radius: 10px;
        padding: 10px;
    }
    div[data-testid="stMarkdownContainer"] p, div[data-testid="stMarkdownContainer"] h1,
    div[data-testid="stMarkdownContainer"] h2, div[data-testid="stMarkdownContainer"] h3,
    div[data-testid="stMarkdownContainer"] h4 {
        color: #FFFFFF !important; text-shadow: 1px 1px 3px rgba(0,0,0,0.8);
    }
    [data-testid="stAlert"] div[data-testid="stMarkdownContainer"] p,
    [data-testid="stAlert"] div[data-testid="stMarkdownContainer"] h3 {
        color: #262730 !important; text-shadow: none !important;
    }
    div[style*="background-color: rgba(255, 255, 255, 0.95)"] {
        color: #262730 !important; text-shadow: none !important;
    }
    label[data-testid="stWidgetLabel"] p {
        color: white !important; font-size: 16px !important; font-weight: 600 !important; text-shadow: 1px 1px 4px rgba(0,0,0,0.6);
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 15px !important; background-color: transparent !important; flex-wrap: wrap !important;
    }
    button[data-baseweb="tab"] div[data-testid="stMarkdownContainer"] p {
        color: #C0C0C0 !important; font-size: 17px !important; font-weight: normal !important; text-shadow: 1px 1px 3px rgba(0,0,0,0.8) !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] div[data-testid="stMarkdownContainer"] p {
        color: #FFFFFF !important; font-weight: bold !important; text-shadow: 1px 1px 5px rgba(0,0,0,0.8) !important;
    }
    div[data-baseweb="tab-highlight"] { background-color: #FFFFFF !important; }
    div[data-baseweb="tab-border"] { display: none !important; background-color: transparent !important; }
    </style>
    """,
    unsafe_allow_html=True
)

st.markdown('<h1 style="color: white; text-shadow: 2px 2px 5px rgba(0,0,0,0.6);">Asistente Operativo - Área Vida</h1>', unsafe_allow_html=True)
st.markdown('<p style="color: white; font-size: 18px; text-shadow: 1px 1px 4px rgba(0,0,0,0.6); margin-bottom: 30px;">Seleccioná la herramienta que necesites usar hoy.</p>', unsafe_allow_html=True)

tab_armador, tab_validador, tab_unificador, tab_pdf = st.tabs([
    "🪄 Armar Excel", "✅ Validar Carga", "📑 Procesar Multi-Hoja", "📄 PDF a Excel / Word"
])

# ==========================================
# PESTAÑA 1: ARMADOR / LIMPIADOR DE DATOS CRUDOS
# ==========================================
with tab_armador:
    st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; margin-bottom: 20px; color: #262730 !important;">Subí el Excel desordenado que te envió la empresa. Podés indicarle al sistema en qué fila empiezan los títulos para saltear logos o membretes.</div>', unsafe_allow_html=True)
    
    col_fila, col_espacio = st.columns([2, 3])
    with col_fila:
        fila_titulos = st.number_input("¿En qué fila del Excel están los títulos de las columnas?", min_value=1, value=1)
    
    col1a, col2a = st.columns([2, 3])
    with col1a:
        archivo_crudo = st.file_uploader("Subí el Excel crudo del cliente", type=["xlsx", "xls", "csv"], key=f"uploader_armador_{st.session_state.reset_armador}")
    
    if archivo_crudo is not None:
        try:
            filas_a_saltar = int(fila_titulos) - 1
            if archivo_crudo.name.endswith('.csv'):
                df_crudo = pd.read_csv(archivo_crudo, skiprows=filas_a_saltar)
            else:
                df_crudo = pd.read_excel(archivo_crudo, skiprows=filas_a_saltar)
            
            df_crudo = df_crudo.dropna(how='all')
            columnas_disponibles = ["No incluir"] + list(df_crudo.columns)
            
            st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; color: #262730 !important;">', unsafe_allow_html=True)
            st.write("### 📌 Mapeo de Columnas")
            
            formato_nombres = st.radio("¿Cómo vienen el Nombre y Apellido en el Excel original?", ["En una sola columna", "En columnas separadas"], horizontal=True)
            
            if formato_nombres == "En una sola columna":
                c1, c2, c3, c4 = st.columns(4)
                with c1: col_dni = st.selectbox("Columna de DNI:", options=columnas_disponibles)
                with c2: col_nombre_completo = st.selectbox("Nombre y Apellido juntos:", options=columnas_disponibles)
                with c3: col_fecha = st.selectbox("Fecha de Nacimiento:", options=columnas_disponibles)
                with c4: col_id_persona = st.selectbox("ID Persona (Opcional):", options=columnas_disponibles)
                col_apellido_sep, col_nombre_sep = "No incluir", "No incluir"
            else:
                c1, c2, c3, c4, c5 = st.columns(5)
                with c1: col_dni = st.selectbox("Columna de DNI:", options=columnas_disponibles)
                with c2: col_apellido_sep = st.selectbox("Columna de Apellido:", options=columnas_disponibles)
                with c3: col_nombre_sep = st.selectbox("Columna de Nombre:", options=columnas_disponibles)
                with c4: col_fecha = st.selectbox("Fecha Nacimiento:", options=columnas_disponibles)
                with c5: col_id_persona = st.selectbox("ID (Opcional):", options=columnas_disponibles)
                col_nombre_completo = "No incluir"
                
            if st.button("Procesar y Generar Plantilla Oficial"):
                columnas_oficiales = ['Acción', '*Tipo Id.', '*Nro. Id.', '*Apellido', '*Nombre', '*Fecha de Nacimiento', 'Fecha de Ingreso', 'Sueldo', 'Suma asegurada', 'Nro. de grupo']
                
                if col_id_persona != "No incluir":
                    columnas_oficiales.insert(1, 'ID de Persona')
                    
                df_oficial = pd.DataFrame(columns=columnas_oficiales)
                
                if col_id_persona != "No incluir":
                    df_oficial['ID de Persona'] = df_crudo[col_id_persona].fillna("")
                
                if col_dni != "No incluir":
                    df_oficial['*Nro. Id.'] = df_crudo[col_dni].apply(procesar_dni_cuil)
                
                apellidos, nombres = [], []
                
                if formato_nombres == "En una sola columna":
                    if col_nombre_completo != "No incluir":
                        for nombre_completo in df_crudo[col_nombre_completo]:
                            texto = str(nombre_completo).strip()
                            if texto.lower() == 'nan' or pd.isna(nombre_completo) or not texto:
                                apellidos.append(""); nombres.append("")
                            elif "," in texto:
                                partes = texto.split(",", 1)
                                apellidos.append(partes[0].strip().upper())
                                nombres.append(partes[1].strip().upper())
                            else:
                                partes = texto.split(" ", 1)
                                if len(partes) > 1:
                                    apellidos.append(partes[0].strip().upper())
                                    nombres.append(partes[1].strip().upper())
                                else:
                                    apellidos.append(texto.upper()); nombres.append("")
                    else:
                        apellidos = [""] * len(df_crudo); nombres = [""] * len(df_crudo)
                else: 
                    if col_apellido_sep != "No incluir":
                        apellidos = [str(x).strip().upper() if pd.notna(x) and str(x).lower() != 'nan' else "" for x in df_crudo[col_apellido_sep]]
                    else:
                        apellidos = [""] * len(df_crudo)
                        
                    if col_nombre_sep != "No incluir":
                        nombres = [str(x).strip().upper() if pd.notna(x) and str(x).lower() != 'nan' else "" for x in df_crudo[col_nombre_sep]]
                    else:
                        nombres = [""] * len(df_crudo)

                df_oficial['*Apellido'] = apellidos
                df_oficial['*Nombre'] = nombres
                
                if col_fecha != "No incluir":
                    fechas_limpias = []
                    for f in df_crudo[col_fecha]:
                        if pd.isna(f) or str(f).lower() == 'nan':
                            fechas_limpias.append("")
                        elif isinstance(f, (pd.Timestamp, datetime)):
                            fechas_limpias.append(f.date())
                        else:
                            f_str = str(f)
                            nums = re.sub(r'[^0-9]', '', f_str)
                            if len(nums) == 8:
                                try: fechas_limpias.append(pd.to_datetime(nums, format='%d%m%Y').date())
                                except: fechas_limpias.append(f_str)
                            else:
                                fechas_limpias.append(f_str)
                    df_oficial['*Fecha de Nacimiento'] = fechas_limpias

                df_oficial['Acción'] = 'Agregar'
                df_oficial['*Tipo Id.'] = 'D.N.I.'
                df_oficial['Nro. de grupo'] = '1-Titular'
                
                dnis_con_datos = df_oficial[df_oficial['*Nro. Id.'] != '']
                cantidad_duplicados = dnis_con_datos.duplicated(subset=['*Nro. Id.', '*Apellido', '*Nombre'], keep='first').sum()

                if cantidad_duplicados > 0:
                    m_vacios = df_oficial['*Nro. Id.'] == ''
                    m_duplicados = df_oficial.duplicated(subset=['*Nro. Id.', '*Apellido', '*Nombre'], keep='first')
                    df_oficial = df_oficial[m_vacios | (~m_duplicados)]
                    df_oficial.reset_index(drop=True, inplace=True)
                    st.success(f"🪄 **Limpieza Automática:** Se eliminaron {cantidad_duplicados} registros idénticos (mismo DNI, Nombre y Apellido).")
                else:
                    st.success("✅ Plantilla generada exitosamente.")
                
                dnis_con_datos_final = df_oficial[df_oficial['*Nro. Id.'] != '']
                cant_dup_parciales = dnis_con_datos_final.duplicated(subset=['*Nro. Id.'], keep=False).sum()

                if cant_dup_parciales > 0:
                    st.warning(f"⚠️ **¡Alerta de revisión!** Quedaron {cant_dup_parciales} registros con el mismo DNI pero distintos nombres.")
                
                st.dataframe(df_oficial.style.apply(pintar_duplicados, subset=['*Nro. Id.']))
                
                buffer_armado = io.BytesIO()
                with pd.ExcelWriter(buffer_armado, engine='xlsxwriter', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
                    df_oficial.to_excel(writer, index=False, sheet_name="VIDA - Personas x Grupo")
                    workbook = writer.book
                    worksheet = writer.sheets['VIDA - Personas x Grupo']
                    inyectar_listas_desplegables(writer, df_oficial, workbook, "VIDA - Personas x Grupo")
                    
                    idx_dni = list(df_oficial.columns).index('*Nro. Id.')
                    letra_col_dni = string.ascii_uppercase[idx_dni]
                    rango_dni = f"{letra_col_dni}2:{letra_col_dni}5000"
                    formula = f'=AND({letra_col_dni}2<>"", COUNTIF(${letra_col_dni}$2:${letra_col_dni}$5000,{letra_col_dni}2)>1)'
                    
                    formato_amarillo = workbook.add_format({'bg_color': '#FFF2CC', 'font_color': '#9C6500'})
                    worksheet.conditional_format(rango_dni, {'type': 'formula', 'criteria': formula, 'format': formato_amarillo})
                
                col_d1, col_d2 = st.columns([1, 1])
                with col_d1:
                    st.download_button(
                        label="📥 Descargar Plantilla Oficial Pre-armada",
                        data=buffer_armado.getvalue(),
                        file_name="VIDA_Personas_Prearmado.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )
                with col_d2:
                    if st.button("🧹 Empezar de cero / Nueva Carga", key="btn_armador", use_container_width=True):
                        st.session_state.reset_armador += 1
                        st.rerun()

            st.markdown('</div>', unsafe_allow_html=True)
        except Exception as e:
            st.error(f"Error al procesar el archivo: {e}. Verificá que la fila de inicio sea la correcta.")

# ==========================================
# PESTAÑA 2: VALIDADOR ESTRICTO
# ==========================================
with tab_validador:
    col1b, col2b = st.columns([2, 3])
    with col1b:
        archivo_agente = st.file_uploader("Seleccioná el archivo Excel listo (.xlsx)", type=["xlsx"], key=f"uploader_validador_{st.session_state.reset_validador}")

    if archivo_agente is not None:
        try:
            # Aquí leemos solo la primera hoja que contenga los datos para la validación tradicional
            hojas = pd.ExcelFile(archivo_agente).sheet_names
            hoja_a_leer = hojas[0] if "DataValidation" not in hojas[0] else hojas[1]
            
            df = pd.read_excel(archivo_agente, sheet_name=hoja_a_leer)
            st.info(f"Procesando validaciones de la hoja '{hoja_a_leer}'...")

            hoy = pd.Timestamp.today()
            total_fechas_mal_formato = 0
            total_fechas_corregidas = 0
            
            errores_col = {
                'Err_TipoId': [''] * len(df), 'Err_NroId': [''] * len(df), 'Err_Apellido': [''] * len(df), 'Err_Nombre': [''] * len(df), 'Err_FechaNac': [''] * len(df),
                'Err_FechaIng': [''] * len(df), 'Err_Sueldo': [''] * len(df), 'Err_SumaAseg': [''] * len(df), 'Err_Grupo': [''] * len(df)
            }

            for index, row in df.iterrows():
                apellido = str(row.get('*Apellido', '')).strip()
                if pd.isna(row.get('*Apellido')) or apellido.lower() == 'nan' or not apellido:
                    errores_col['Err_Apellido'][index] = "Apellido vacío"
                    
                nombre = str(row.get('*Nombre', '')).strip()
                if pd.isna(row.get('*Nombre')) or nombre.lower() == 'nan' or not nombre:
                    errores_col['Err_Nombre'][index] = "Nombre vacío"

                tipo_id = str(row.get('*Tipo Id.', '')).strip().upper()
                if pd.isna(row.get('*Tipo Id.')) or tipo_id == 'NAN' or not tipo_id:
                    errores_col['Err_TipoId'][index] = "Falta seleccionar D.N.I. o Pasaporte"
                
                nro_id = row.get('*Nro. Id.')
                nro_id_str = str(nro_id).strip()
                
                if pd.isna(nro_id) or nro_id_str == 'nan' or not nro_id_str:
                    errores_col['Err_NroId'][index] = "DNI vacío"
                else:
                    if not nro_id_str.replace('.', '').isdigit() or '.' in nro_id_str or ',' in nro_id_str:
                        errores_col['Err_NroId'][index] = "DNI con letras/puntos"
                    elif tipo_id == 'PASAPORTE':
                        if not nro_id_str.isdigit() or nro_id_str.startswith('0'):
                            errores_col['Err_NroId'][index] = "Pasaporte inicia con 0 o tiene letras"
                
                # Validar Fecha de Nacimiento
                fecha_nac = row.get('*Fecha de Nacimiento')
                if pd.notna(fecha_nac):
                    es_valida = True
                    if not isinstance(fecha_nac, (pd.Timestamp, datetime)):
                        total_fechas_mal_formato += 1
                        fecha_str = str(fecha_nac)
                        numeros = re.sub(r'[^0-9]', '', fecha_str)
                        corregida = False
                        
                        if len(numeros) == 8:
                            try:
                                fecha_limpia = pd.to_datetime(numeros, format='%d%m%Y')
                                df.at[index, '*Fecha de Nacimiento'] = fecha_limpia
                                fecha_nac = fecha_limpia
                                corregida = True
                                total_fechas_corregidas += 1
                            except: pass
                        
                        if not corregida:
                            errores_col['Err_FechaNac'][index] = "Fecha Nac. irreconocible"
                            es_valida = False

                    if es_valida and (isinstance(fecha_nac, pd.Timestamp) or isinstance(fecha_nac, datetime)):
                        if fecha_nac > hoy:
                            errores_col['Err_FechaNac'][index] = "Fecha futura"
                        else:
                            edad = hoy.year - fecha_nac.year - ((hoy.month, hoy.day) < (fecha_nac.month, fecha_nac.day))
                            if edad > 80: errores_col['Err_FechaNac'][index] = "Mayor de 80 años"
                            elif edad < 14: errores_col['Err_FechaNac'][index] = "Menor de 14 años"

                # Validar Fecha de Ingreso
                fecha_ing = row.get('Fecha de Ingreso')
                if pd.notna(fecha_ing) and str(fecha_ing).strip() != '' and str(fecha_ing).lower() != 'nan':
                    if not isinstance(fecha_ing, (pd.Timestamp, datetime)):
                        errores_col['Err_FechaIng'][index] = "Fecha Ingreso formato inválido"
                    elif isinstance(fecha_ing, (pd.Timestamp, datetime)) and fecha_ing > hoy:
                        errores_col['Err_FechaIng'][index] = "Fecha Ingreso futura"

                # Validar Sueldo y Suma asegurada
                for col_moneda, err_key in [('Sueldo', 'Err_Sueldo'), ('Suma asegurada', 'Err_SumaAseg')]:
                    val = row.get(col_moneda)
                    if pd.notna(val) and str(val).strip() != '' and str(val).lower() != 'nan':
                        try: float(val)
                        except: errores_col[err_key][index] = f"{col_moneda} no es número"

                # Validar Nro de grupo
                grupo = str(row.get('Nro. de grupo', '')).strip()
                if grupo and grupo != 'nan' and grupo not in LISTAS_DESPLEGABLES['Nro. de grupo']:
                    errores_col['Err_Grupo'][index] = "Nro. de grupo inválido"

            nro_ids = df['*Nro. Id.'].dropna().astype(str).tolist()
            dups = set([x for x in nro_ids if nro_ids.count(x) > 1])

            for index, row in df.iterrows():
                nro_id = str(row.get('*Nro. Id.'))
                if nro_id in dups:
                    actual = errores_col['Err_NroId'][index]
                    errores_col['Err_NroId'][index] = "DNI duplicado" if not actual else actual + " / DNI duplicado"

            lista_todos_errores = []
            for categoria, lista_err in errores_col.items():
                for error in lista_err:
                    if error != '':
                        for sub_error in error.split(" / "):
                            lista_todos_errores.append(sub_error)
            
            conteo_errores = Counter(lista_todos_errores)
            tiene_error_fila = [False] * len(df)
            
            mapa_nombres = {
                'Err_TipoId': 'Error: Tipo Id', 'Err_NroId': 'Error: Nro Id', 'Err_Apellido': 'Error: Apellido', 
                'Err_Nombre': 'Error: Nombre', 'Err_FechaNac': 'Error: Fecha Nac', 'Err_FechaIng': 'Error: Fecha Ingreso',
                'Err_Sueldo': 'Error: Sueldo', 'Err_SumaAseg': 'Error: Suma Asegurada', 'Err_Grupo': 'Error: Grupo'
            }

            columnas_con_errores = []
            for clave, nombre_col in mapa_nombres.items():
                if any(errores_col[clave]): 
                    df[nombre_col] = errores_col[clave]
                    columnas_con_errores.append(nombre_col)
                    for i, val in enumerate(errores_col[clave]):
                        if val != '': tiene_error_fila[i] = True

            df.insert(0, 'Estado Fila', ["❌ ERROR" if e else "✅ OK" for e in tiene_error_fila])
            errores_totales = sum(tiene_error_fila)
            
            if total_fechas_mal_formato > 0:
                st.success(f"🪄 **Autocorrección Inteligente:** Se detectaron {total_fechas_mal_formato} fechas mal escritas. El sistema logró corregir automáticamente {total_fechas_corregidas} de ellas.")
            
            if errores_totales == 0:
                st.success("✅ ¡Excelente! El archivo ya no tiene errores y está listo para enviar.")
            else:
                st.error(f"❌ Se encontraron {errores_totales} filas con errores que requieren intervención manual.")
                
                html_resumen = f"""
                <div style="background-color: rgba(255, 255, 255, 0.95); padding: 20px; border-radius: 10px; color: #333; margin-bottom: 20px; box-shadow: 0 4px 6px rgba(0,0,0,0.1);">
                    <h3 style="color: #d32f2f; margin-top: 0;">📊 Detalle de Errores Restantes:</h3>
                    <ul style="font-size: 16px;">
                """
                for error_texto, cantidad in conteo_errores.most_common():
                    html_resumen += f"<li><b>{cantidad}</b> x {error_texto}</li>"
                    
                html_resumen += """
                    </ul>
                    <hr style="border-top: 1px solid #ccc;">
                    <p style="margin-bottom: 0;"><b>Vista previa</b> (descargá el Excel para ver todos los detalles y usar los filtros):</p>
                </div>
                """
                st.markdown(html_resumen, unsafe_allow_html=True)
                st.dataframe(df[df['Estado Fila'] == "❌ ERROR"][['Estado Fila', '*Nro. Id.', '*Apellido'] + columnas_con_errores])

            df['*Fecha de Nacimiento'] = df['*Fecha de Nacimiento'].apply(lambda x: x.date() if isinstance(x, (pd.Timestamp, datetime)) else x)
            if 'Fecha de Ingreso' in df.columns:
                df['Fecha de Ingreso'] = df['Fecha de Ingreso'].apply(lambda x: x.date() if isinstance(x, (pd.Timestamp, datetime)) else x)
            
            buffer = io.BytesIO()
            with pd.ExcelWriter(buffer, engine='xlsxwriter', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
                df.to_excel(writer, index=False, sheet_name=hoja_a_leer)
                workbook = writer.book
                
                inyectar_listas_desplegables(writer, df, workbook, hoja_a_leer)
                
                formato_rojo = workbook.add_format({'bg_color': '#FFC7CE', 'font_color': '#9C0006'})
                formato_verde = workbook.add_format({'bg_color': '#C6EFCE', 'font_color': '#006100'})
                formato_amarillo = workbook.add_format({'bg_color': '#FFF2CC', 'font_color': '#9C6500'})
                worksheet = writer.sheets[hoja_a_leer]
                
                worksheet.conditional_format('A2:A5000', {'type': 'text', 'criteria': 'containing', 'value': 'ERROR', 'format': formato_rojo})
                worksheet.conditional_format('A2:A5000', {'type': 'text', 'criteria': 'containing', 'value': 'OK', 'format': formato_verde})
                worksheet.set_column(0, 0, 15) 
                
                if columnas_con_errores:
                    idx_inicio = len(df.columns) - len(columnas_con_errores)
                    worksheet.set_column(idx_inicio, len(df.columns)-1, 25, formato_amarillo)

            col_d1, col_d2 = st.columns([1, 1])
            with col_d1:
                st.download_button(
                    label="📥 Descargar Excel con Reporte de Errores" if errores_totales > 0 else "📥 Descargar Excel Validado",
                    data=buffer.getvalue(),
                    file_name="Control_VIDA_Reporte.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            with col_d2:
                if st.button("🧹 Empezar de cero / Nueva Carga", key="btn_validador", use_container_width=True):
                    st.session_state.reset_validador += 1
                    st.rerun()

        except Exception as e:
            st.error(f"Error al leer el archivo. Detalle técnico: {e}")

# ==========================================
# PESTAÑA 3: PROCESADOR MULTI-HOJA
# ==========================================
with tab_unificador:
    st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; margin-bottom: 20px; color: #262730 !important;">Subí un archivo Excel con varias hojas (o varios archivos). Podés configurar cada hoja por separado y el sistema generará un <b>único Excel de salida respetando las hojas separadas</b>, aplicando los controles de armado, fechas de nacimiento y eliminando duplicados <b>solo dentro de cada hoja</b>.</div>', unsafe_allow_html=True)
    
    archivos_multiples = st.file_uploader("Subí los archivos que quieras procesar", type=["xlsx", "xls", "csv"], accept_multiple_files=True, key=f"uploader_unif_{st.session_state.reset_unificador}")
    
    if archivos_multiples:
        configs = []
        st.write("### ⚙️ Configuración de Archivos y Hojas")
        
        idx_global = 0
        for archivo in archivos_multiples:
            es_excel = not archivo.name.endswith('.csv')
            hojas = ["CSV"]
            
            if es_excel:
                archivo.seek(0)
                try:
                    xls = pd.ExcelFile(archivo)
                    hojas = xls.sheet_names
                except:
                    hojas = ["Hoja 1"]
            
            for hoja in hojas:
                with st.expander(f"📁 {archivo.name} - Hoja: {hoja}", expanded=True):
                    incluir_hoja = st.checkbox(f"Procesar esta hoja", value=True, key=f"check_{idx_global}")
                    
                    if incluir_hoja:
                        col_f_1, col_f_2 = st.columns([2, 3])
                        with col_f_1:
                            fila_titulos = st.number_input("Fila de títulos:", min_value=1, value=2, key=f"fila_unif_{idx_global}")
                        
                        try:
                            filas_a_saltar = int(fila_titulos) - 1
                            archivo.seek(0)
                            
                            if es_excel:
                                df_tmp = pd.read_excel(archivo, sheet_name=hoja, skiprows=filas_a_saltar, nrows=5)
                            else:
                                df_tmp = pd.read_csv(archivo, skiprows=filas_a_saltar, nrows=5)
                                
                            columnas_disponibles = ["No incluir"] + list(df_tmp.columns)
                            
                            formato_nombres = st.radio("¿Cómo vienen el Nombre y Apellido acá?", ["En una sola columna", "En columnas separadas"], horizontal=True, key=f"fmt_nom_unif_{idx_global}")
                            
                            if formato_nombres == "En una sola columna":
                                c1, c2, c3, c4 = st.columns(4)
                                with c1: col_dni = st.selectbox("DNI:", options=columnas_disponibles, key=f"dni1_{idx_global}")
                                with c2: col_nombre_completo = st.selectbox("Nombre completo:", options=columnas_disponibles, key=f"nomc_{idx_global}")
                                with c3: col_fecha = st.selectbox("Fecha Nac.:", options=columnas_disponibles, key=f"fec1_{idx_global}")
                                with c4: col_id_persona = st.selectbox("ID (Opcional):", options=columnas_disponibles, key=f"id1_{idx_global}")
                                col_apellido_sep, col_nombre_sep = "No incluir", "No incluir"
                            else:
                                c1, c2, c3, c4, c5 = st.columns(5)
                                with c1: col_dni = st.selectbox("DNI:", options=columnas_disponibles, key=f"dni2_{idx_global}")
                                with c2: col_apellido_sep = st.selectbox("Apellido:", options=columnas_disponibles, key=f"ape_{idx_global}")
                                with c3: col_nombre_sep = st.selectbox("Nombre:", options=columnas_disponibles, key=f"nom_{idx_global}")
                                with c4: col_fecha = st.selectbox("Fecha Nac.:", options=columnas_disponibles, key=f"fec2_{idx_global}")
                                with c5: col_id_persona = st.selectbox("ID (Opcional):", options=columnas_disponibles, key=f"id2_{idx_global}")
                                col_nombre_completo = "No incluir"
                                
                            configs.append({
                                'archivo': archivo, 'es_excel': es_excel, 'hoja': hoja, 'fila_titulos': fila_titulos, 
                                'formato_nombres': formato_nombres, 'col_dni': col_dni, 'col_nombre_completo': col_nombre_completo, 
                                'col_apellido_sep': col_apellido_sep, 'col_nombre_sep': col_nombre_sep, 'col_fecha': col_fecha, 'col_id_persona': col_id_persona
                            })
                        except Exception as e:
                            st.warning(f"No se pudo leer la tabla en esta hoja. Detalle: {e}")
                    
                idx_global += 1

        if configs:
            st.markdown("---")
            if st.button("📑 Procesar Hojas y Generar Plantilla", use_container_width=True):
                hojas_procesadas = []
                tiene_id = any(cfg['col_id_persona'] != "No incluir" for cfg in configs)
                
                with st.spinner("Procesando hojas..."):
                    for cfg in configs:
                        cfg['archivo'].seek(0)
                        filas_a_saltar = int(cfg['fila_titulos']) - 1
                        
                        if cfg['es_excel']:
                            df_crudo = pd.read_excel(cfg['archivo'], sheet_name=cfg['hoja'], skiprows=filas_a_saltar)
                        else:
                            df_crudo = pd.read_csv(cfg['archivo'], skiprows=filas_a_saltar)
                        
                        df_crudo = df_crudo.dropna(how='all')
                        
                        columnas_oficiales = ['Acción', '*Tipo Id.', '*Nro. Id.', '*Apellido', '*Nombre', '*Fecha de Nacimiento', 'Fecha de Ingreso', 'Sueldo', 'Suma asegurada', 'Nro. de grupo']
                        if tiene_id:
                            columnas_oficiales.insert(1, 'ID de Persona')
                            
                        df_oficial = pd.DataFrame(columns=columnas_oficiales)
                        
                        if tiene_id:
                            if cfg['col_id_persona'] != "No incluir":
                                df_oficial['ID de Persona'] = df_crudo[cfg['col_id_persona']].fillna("")
                            else:
                                df_oficial['ID de Persona'] = ""
                                
                        if cfg['col_dni'] != "No incluir":
                            df_oficial['*Nro. Id.'] = df_crudo[cfg['col_dni']].apply(procesar_dni_cuil)
                        
                        apellidos, nombres = [], []
                        if cfg['formato_nombres'] == "En una sola columna":
                            if cfg['col_nombre_completo'] != "No incluir":
                                for nombre_completo in df_crudo[cfg['col_nombre_completo']]:
                                    texto = str(nombre_completo).strip()
                                    if texto.lower() == 'nan' or pd.isna(nombre_completo) or not texto:
                                        apellidos.append(""); nombres.append("")
                                    elif "," in texto:
                                        partes = texto.split(",", 1)
                                        apellidos.append(partes[0].strip().upper())
                                        nombres.append(partes[1].strip().upper())
                                    else:
                                        partes = texto.split(" ", 1)
                                        if len(partes) > 1:
                                            apellidos.append(partes[0].strip().upper())
                                            nombres.append(partes[1].strip().upper())
                                        else:
                                            apellidos.append(texto.upper()); nombres.append("")
                            else:
                                apellidos = [""] * len(df_crudo); nombres = [""] * len(df_crudo)
                        else:
                            if cfg['col_apellido_sep'] != "No incluir":
                                apellidos = [str(x).strip().upper() if pd.notna(x) and str(x).lower() != 'nan' else "" for x in df_crudo[cfg['col_apellido_sep']]]
                            else:
                                apellidos = [""] * len(df_crudo)
                                
                            if cfg['col_nombre_sep'] != "No incluir":
                                nombres = [str(x).strip().upper() if pd.notna(x) and str(x).lower() != 'nan' else "" for x in df_crudo[cfg['col_nombre_sep']]]
                            else:
                                nombres = [""] * len(df_crudo)

                        df_oficial['*Apellido'] = apellidos
                        df_oficial['*Nombre'] = nombres
                        
                        if cfg['col_fecha'] != "No incluir":
                            fechas_limpias = []
                            for f in df_crudo[cfg['col_fecha']]:
                                if pd.isna(f) or str(f).lower() == 'nan':
                                    fechas_limpias.append("")
                                elif isinstance(f, (pd.Timestamp, datetime)):
                                    fechas_limpias.append(f.date())
                                else:
                                    f_str = str(f)
                                    nums = re.sub(r'[^0-9]', '', f_str)
                                    if len(nums) == 8:
                                        try: fechas_limpias.append(pd.to_datetime(nums, format='%d%m%Y').date())
                                        except: fechas_limpias.append(f_str)
                                    else:
                                        fechas_limpias.append(f_str)
                            df_oficial['*Fecha de Nacimiento'] = fechas_limpias

                        df_oficial['Acción'] = 'Agregar'
                        df_oficial['*Tipo Id.'] = 'D.N.I.'
                        df_oficial['Nro. de grupo'] = '1-Titular'
                        
                        # --- LIMPIEZA DE DUPLICADOS EXACTOS (SOLO DENTRO DE LA HOJA) ---
                        dnis_con_datos = df_oficial[df_oficial['*Nro. Id.'] != '']
                        cantidad_duplicados = dnis_con_datos.duplicated(subset=['*Nro. Id.', '*Apellido', '*Nombre'], keep='first').sum()

                        if cantidad_duplicados > 0:
                            m_vacios = df_oficial['*Nro. Id.'] == ''
                            m_duplicados = df_oficial.duplicated(subset=['*Nro. Id.', '*Apellido', '*Nombre'], keep='first')
                            df_oficial = df_oficial[m_vacios | (~m_duplicados)]
                            df_oficial.reset_index(drop=True, inplace=True)
                            
                        hojas_procesadas.append({
                            'nombre_hoja': cfg['hoja'],
                            'df': df_oficial,
                            'duplicados_eliminados': cantidad_duplicados
                        })

                st.success("✅ ¡Proceso completado! Se generó el Excel conservando las hojas separadas y eliminando duplicados internos.")
                
                # Resumen visual por hoja
                st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; color: #262730 !important;">', unsafe_allow_html=True)
                for item in hojas_procesadas:
                    st.write(f"**Hoja: {item['nombre_hoja']}**")
                    if item['duplicados_eliminados'] > 0:
                        st.info(f"🪄 Se eliminaron {item['duplicados_eliminados']} registros idénticos dentro de esta hoja.")
                    st.dataframe(item['df'].head(5)) # Vista previa cortita para no saturar
                st.markdown('</div>', unsafe_allow_html=True)
                
                # Generación del archivo Multi-Hoja
                buffer_multi = io.BytesIO()
                with pd.ExcelWriter(buffer_multi, engine='xlsxwriter', datetime_format='dd/mm/yyyy', date_format='dd/mm/yyyy') as writer:
                    workbook = writer.book
                    
                    for idx, item in enumerate(hojas_procesadas):
                        # Limpiamos el nombre de la hoja para que Excel lo acepte (máx 31 caracteres)
                        nombre_limpio = re.sub(r'[\[\]\:\*\?\/\\]', '', str(item['nombre_hoja']))[:31]
                        
                        # Si casualmente dos hojas quedan con el mismo nombre tras limpiarlo, evitamos error
                        if nombre_limpio in writer.sheets:
                            nombre_limpio = f"{idx+1}_{nombre_limpio}"[:31]
                            
                        item['df'].to_excel(writer, index=False, sheet_name=nombre_limpio)
                        worksheet = writer.sheets[nombre_limpio]
                        
                        # Inyectamos listas desplegables hoja por hoja
                        inyectar_listas_desplegables(writer, item['df'], workbook, nombre_limpio)
                        
                        # Inyectamos el formato condicional amarillo (para duplicados parciales DENTRO de la hoja)
                        idx_dni = list(item['df'].columns).index('*Nro. Id.')
                        letra_col_dni = string.ascii_uppercase[idx_dni]
                        rango_dni = f"{letra_col_dni}2:{letra_col_dni}5000"
                        formula = f'=AND({letra_col_dni}2<>"", COUNTIF(${letra_col_dni}$2:${letra_col_dni}$5000,{letra_col_dni}2)>1)'
                        
                        formato_amarillo = workbook.add_format({'bg_color': '#FFF2CC', 'font_color': '#9C6500'})
                        worksheet.conditional_format(rango_dni, {'type': 'formula', 'criteria': formula, 'format': formato_amarillo})
                
                col_u1, col_u2 = st.columns([1, 1])
                with col_u1:
                    st.download_button(
                        label="📥 Descargar Excel Multi-Hoja Procesado",
                        data=buffer_multi.getvalue(),
                        file_name="VIDA_Personas_MultiHoja.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )
                with col_u2:
                    if st.button("🧹 Empezar de cero / Nueva Carga", key="btn_unif", use_container_width=True):
                        st.session_state.reset_unificador += 1
                        st.rerun()

# ==========================================
# PESTAÑA 4: CONVERTIDOR DE PDF
# ==========================================
with tab_pdf:
    st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; margin-bottom: 20px; color: #262730 !important;">Subí la nómina en PDF. Si el sistema detecta una tabla estructurada, generará un <b>Excel</b>. Si solo detecta texto suelto, generará un archivo de <b>Word</b> para que te sea más fácil copiar y pegar los datos.</div>', unsafe_allow_html=True)
    
    col1c, col2c = st.columns([2, 3])
    with col1c:
        archivo_pdf = st.file_uploader("Subí el PDF del cliente", type=["pdf"], key=f"uploader_pdf_{st.session_state.reset_pdf}")
    
    if archivo_pdf is not None:
        if st.button("🔍 Extraer Datos"):
            try:
                import pdfplumber
                import docx
            except ImportError:
                st.error("❌ Faltan instalar librerías. Por favor, agregá las palabras 'pdfplumber' y 'python-docx' a tu archivo requirements.txt en GitHub (una debajo de la otra) y esperá a que la aplicación se reinicie.")
            else:
                with st.spinner("Escaneando PDF... esto puede demorar unos segundos."):
                    todas_las_filas = []
                    texto_crudo = []
                    
                    try:
                        with pdfplumber.open(archivo_pdf) as pdf:
                            for page in pdf.pages:
                                tablas = page.extract_tables()
                                if tablas:
                                    for tabla in tablas:
                                        tabla_limpia = [[str(celda).strip() if celda is not None else "" for celda in fila] for fila in tabla]
                                        todas_las_filas.extend(tabla_limpia)
                                
                                texto = page.extract_text()
                                if texto:
                                    texto_crudo.extend(texto.split('\n'))
                        
                        st.markdown('<div style="background-color: rgba(255, 255, 255, 0.95); padding: 15px; border-radius: 10px; color: #262730 !important;">', unsafe_allow_html=True)
                        
                        if todas_las_filas:
                            st.success("✅ ¡Se detectaron tablas estructuradas! Generando archivo Excel...")
                            df_pdf = pd.DataFrame(todas_las_filas)
                            st.write("**Vista Previa:**")
                            st.dataframe(df_pdf.head(15))
                            
                            buffer_pdf = io.BytesIO()
                            with pd.ExcelWriter(buffer_pdf, engine='xlsxwriter') as writer:
                                df_pdf.to_excel(writer, index=False, header=False, sheet_name="Datos Extraídos")
                                workbook = writer.book
                                inyectar_listas_desplegables(writer, df_pdf, workbook, "Datos Extraídos")
                            
                            col_d1, col_d2 = st.columns([1, 1])
                            with col_d1:
                                st.download_button(
                                    label="📥 Descargar Tabla en Excel",
                                    data=buffer_pdf.getvalue(),
                                    file_name="PDF_Convertido.xlsx",
                                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                    use_container_width=True
                                )
                            with col_d2:
                                if st.button("🧹 Empezar de cero", key="btn_pdf_excel", use_container_width=True):
                                    st.session_state.reset_pdf += 1
                                    st.rerun()
                            
                        elif texto_crudo:
                            st.warning("⚠️ El PDF no tenía formato de tabla. Se generó un archivo Word con el texto limpio para que puedas copiar y pegar.")
                            
                            doc = docx.Document()
                            doc.add_heading('Texto extraído del PDF', 0)
                            for linea in texto_crudo:
                                if linea.strip(): 
                                    doc.add_paragraph(linea.strip())
                                    
                            buffer_word = io.BytesIO()
                            doc.save(buffer_word)
                            
                            col_d1, col_d2 = st.columns([1, 1])
                            with col_d1:
                                st.download_button(
                                    label="📝 Descargar Texto en Word (.docx)",
                                    data=buffer_word.getvalue(),
                                    file_name="PDF_Texto_Extraido.docx",
                                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                                    use_container_width=True
                                )
                            with col_d2:
                                if st.button("🧹 Empezar de cero", key="btn_pdf_word", use_container_width=True):
                                    st.session_state.reset_pdf += 1
                                    st.rerun()
                        else:
                            st.error("No se detectaron textos ni tablas. Es probable que el PDF sea una foto o imagen escaneada.")
                            
                        st.markdown('</div>', unsafe_allow_html=True)
                        
                    except Exception as e:
                        st.error(f"Ocurrió un error inesperado al leer el PDF: {e}")
